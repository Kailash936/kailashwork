"""
Excel Writer Module
Handles writing processed data to Excel files
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict
import os
from datetime import datetime

class ExcelWriter:
    def __init__(self, output_path: str):
        self.output_path = output_path
        self.workbook = None
        self.worksheet = None
    
    def create_excel_file(self, data: List[Dict[str, str]], template_path: str = None):
        """Create Excel file with processed data"""
        try:
            # Create DataFrame from processed data
            df = pd.DataFrame(data)
            
            # If template exists, use it as base
            if template_path and os.path.exists(template_path):
                # Load existing workbook
                self.workbook = openpyxl.load_workbook(template_path)
                self.worksheet = self.workbook.active
                
                # Clear existing data (except headers)
                for row in range(2, self.worksheet.max_row + 1):
                    for col in range(1, self.worksheet.max_column + 1):
                        self.worksheet.cell(row=row, column=col).value = None
            else:
                # Create new workbook
                self.workbook = openpyxl.Workbook()
                self.worksheet = self.workbook.active
                self.worksheet.title = "Customer Data"
            
            # Write headers if not using template
            if not template_path or not os.path.exists(template_path):
                headers = list(df.columns)
                for col, header in enumerate(headers, 1):
                    cell = self.worksheet.cell(row=1, column=col)
                    cell.value = header
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center")
            
            # Write data
            for row_idx, record in enumerate(data, 2):  # Start from row 2 (after header)
                for col_idx, (key, value) in enumerate(record.items(), 1):
                    cell = self.worksheet.cell(row=row_idx, column=col_idx)
                    cell.value = value
                    cell.alignment = Alignment(horizontal="left")
            
            # Auto-adjust column widths
            self.auto_adjust_columns()
            
            # Add borders
            self.add_borders()
            
            # Save the file
            self.workbook.save(self.output_path)
            print(f"Excel file saved successfully: {self.output_path}")
            
            return True
            
        except Exception as e:
            print(f"Error creating Excel file: {e}")
            return False
    
    def auto_adjust_columns(self):
        """Auto-adjust column widths based on content"""
        for column in self.worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            # Set minimum width and add some padding
            adjusted_width = min(max(max_length + 2, 10), 50)
            self.worksheet.column_dimensions[column_letter].width = adjusted_width
    
    def add_borders(self):
        """Add borders to the data area"""
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Apply borders to all data cells
        for row in self.worksheet.iter_rows():
            for cell in row:
                cell.border = thin_border
    
    def create_summary_sheet(self, validation_report: Dict[str, any]):
        """Create a summary sheet with processing statistics"""
        try:
            # Create new sheet for summary
            summary_sheet = self.workbook.create_sheet("Summary")
            
            # Add summary information
            summary_data = [
                ["Auto-Typing Model Summary", ""],
                ["Generated on:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
                ["", ""],
                ["Processing Statistics", ""],
                ["Total Records Processed:", validation_report['total_records']],
                ["Validation Errors:", validation_report['error_count']],
                ["Success Rate:", f"{validation_report['success_rate']:.1f}%"],
                ["", ""],
                ["Column Information", ""],
                ["CUSTOMER NAME", "Customer full name"],
                ["CRN", "Customer Reference Number (10 digits)"],
                ["PROGRAM", "Program information (to be filled)"],
                ["RO_NAME", "Regional Office Name (to be filled)"],
                ["MANUFACTURERDESC", "Manufacturer Description (to be filled)"],
                ["ASSETCAT", "Asset Category (to be filled)"],
                ["MAKE", "Product Make (to be filled)"],
                ["RO_NAME.1", "Regional Office Name 2 (to be filled)"],
                ["ASSESTCAT", "Asset Category 2 (to be filled)"],
                ["MANUFACTURER", "Manufacturer (to be filled)"]
            ]
            
            # Write summary data
            for row_idx, (label, value) in enumerate(summary_data, 1):
                summary_sheet.cell(row=row_idx, column=1, value=label)
                summary_sheet.cell(row=row_idx, column=2, value=value)
                
                # Style the header row
                if row_idx == 1:
                    summary_sheet.cell(row=row_idx, column=1).font = Font(bold=True, size=14)
                    summary_sheet.cell(row=row_idx, column=2).font = Font(bold=True, size=14)
            
            # Auto-adjust column widths for summary
            for column in summary_sheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max(max_length + 2, 10), 50)
                summary_sheet.column_dimensions[column_letter].width = adjusted_width
            
            return True
            
        except Exception as e:
            print(f"Error creating summary sheet: {e}")
            return False

def main():
    """Test the Excel writer"""
    # Sample data for testing
    sample_data = [
        {'CUSTOMER NAME': 'AANCHAL JOSHI', 'CRN': '5041885433', 'PROGRAM': '', 'RO_NAME': '', 'MANUFACTURERDESC': '', 'ASSETCAT': '', 'MAKE': '', 'RO_NAME.1': '', 'ASSESTCAT': '', 'MANUFACTURER': ''},
        {'CUSTOMER NAME': 'AIMAN ALI', 'CRN': '5111926597', 'PROGRAM': '', 'RO_NAME': '', 'MANUFACTURERDESC': '', 'ASSETCAT': '', 'MAKE': '', 'RO_NAME.1': '', 'ASSESTCAT': '', 'MANUFACTURER': ''}
    ]
    
    validation_report = {
        'total_records': 2,
        'error_count': 0,
        'success_rate': 100.0
    }
    
    writer = ExcelWriter('/workspace/test_output.xlsx')
    success = writer.create_excel_file(sample_data)
    
    if success:
        writer.create_summary_sheet(validation_report)
        writer.workbook.save('/workspace/test_output.xlsx')
        print("Test Excel file created successfully!")

if __name__ == "__main__":
    main()