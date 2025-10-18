#!/usr/bin/env python3

import openpyxl
from openpyxl import load_workbook

def verify_excel_content(excel_path):
    """Verify the content of the updated Excel file"""
    try:
        workbook = load_workbook(excel_path)
        sheet = workbook.active
        
        print(f"=== Excel File Verification: {excel_path} ===")
        print(f"Sheet name: {sheet.title}")
        print(f"Total rows: {sheet.max_row}")
        print(f"Total columns: {sheet.max_column}")
        
        print("\nHeaders:")
        headers = []
        for col in range(1, sheet.max_column + 1):
            header = sheet.cell(row=1, column=col).value
            headers.append(header)
        print(headers)
        
        print(f"\nFirst 15 rows of data:")
        for row in range(1, min(16, sheet.max_row + 1)):
            row_data = []
            for col in range(1, min(6, sheet.max_column + 1)):  # Show first 5 columns
                cell_value = sheet.cell(row=row, column=col).value
                row_data.append(str(cell_value) if cell_value is not None else "")
            print(f"Row {row}: {row_data}")
        
        print(f"\nTotal customers (excluding header): {sheet.max_row - 1}")
        
    except Exception as e:
        print(f"Error reading Excel file: {e}")

if __name__ == "__main__":
    verify_excel_content("/workspace/DOC-20221124-WA0102-2_updated.xlsx")